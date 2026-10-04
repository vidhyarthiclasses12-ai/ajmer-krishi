import json, os, random, sqlite3, uuid
import urllib.parse
import urllib.request
from datetime import datetime, timedelta
from functools import wraps
from io import BytesIO

from flask import Flask, jsonify, request, send_from_directory, session
from openpyxl import Workbook, load_workbook
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.environ.get("DATA_DIR") or ("/data" if os.path.isdir("/data") else os.path.join(BASE, "data"))
DB = os.path.join(DATA, "portal.db")
UPLOAD = os.path.join(DATA, "photos")
os.makedirs(UPLOAD, exist_ok=True)
os.makedirs(os.path.dirname(DB), exist_ok=True)

app = Flask(__name__, static_folder="static")
app.secret_key = os.environ.get("SECRET_KEY", "change-this-ajmer-krishi-secret")
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax", PERMANENT_SESSION_LIFETIME=60*60*8)
FAILS = {}

@app.after_request
def secure_headers(resp):
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "same-origin"
    return resp
app.config["MAX_CONTENT_LENGTH"] = 12 * 1024 * 1024

BLOCKS = {
    "अजमेर ग्रामीण": ["गेगल", "घूघरा", "दौराई", "पुष्कर", "लोहागल", "हटूंडी", "सेदरिया"],
    "बड़ल्या": ["बड़ल्या", "सराधना", "नारेली", "मगरा", "खोड़ा", "आंबा मसीना"],
    "पीसांगन": ["पीसांगन", "गोला", "केसरपुरा", "बुधवाड़ा", "जेठाना", "गणहेड़ा"],
    "किशनगढ़ (सिलोरा)": ["सिलोरा", "रूपनगढ़", "हरमाड़ा", "बांदरसिंदरी", "पाटन", "मोतीपुरा"],
    "नसीराबाद": ["नसीराबाद", "तबाजी", "माकरवाली", "रामसर", "सरसुंडा"],
    "श्रीनगर": ["श्रीनगर", "रामगढ़", "खारवा", "देलवाड़ा"],
    "अराई": ["अराई", "बगहेरा", "जूनिया", "घाटियाली", "देवगांव"],
    "भिनाय": ["भिनाय", "बांदनवाड़ा", "बारली", "देवलिया कलां", "हीरापुरा"],
    "केकड़ी": ["केकड़ी", "जूनिया", "घाटियाली", "धूंधरी", "गुल्गांव"],
    "सरवाड़": ["सरवाड़", "टांटुटी", "रामगढ़", "सांकलिया"],
    "सावर": ["सावर", "रामगढ़", "खारवा", "टीटड़िया"],
}

def db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con

def init():
    con = db()
    con.executescript("""
    CREATE TABLE IF NOT EXISTS users (
      id TEXT PRIMARY KEY, role TEXT, name TEXT, cluster TEXT, mobile TEXT,
      username TEXT UNIQUE, password_hash TEXT, active INTEGER DEFAULT 1
    );
    CREATE TABLE IF NOT EXISTS records (
      id TEXT PRIMARY KEY, user_id TEXT, by_name TEXT, role TEXT, cluster TEXT,
      village TEXT, form_type TEXT, payload TEXT, status TEXT, created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS photos (
      id TEXT PRIMARY KEY, record_id TEXT, filename TEXT, created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS history (
      id TEXT PRIMARY KEY, record_id TEXT, status TEXT, note TEXT, by_name TEXT, at TEXT
    );
    CREATE TABLE IF NOT EXISTS resets (
      id TEXT PRIMARY KEY, username TEXT, status TEXT, created_at TEXT
    );
    CREATE TABLE IF NOT EXISTS clusters (
      id TEXT PRIMARY KEY, district TEXT, block TEXT, name TEXT UNIQUE
    );
    CREATE TABLE IF NOT EXISTS places (
      id TEXT PRIMARY KEY, cluster_id TEXT, village TEXT, gram_panchayat TEXT
    );
    CREATE TABLE IF NOT EXISTS otps (
      id TEXT PRIMARY KEY, username TEXT, mobile TEXT, code TEXT, expires TEXT, used INTEGER DEFAULT 0
    );
    """)
    if not con.execute("SELECT 1 FROM users WHERE username='admin'").fetchone():
        con.execute(
            "INSERT INTO users VALUES (?,?,?,?,?,?,?,1)",
            ("u-admin", "admin", "System Admin", "All", "9000000000", "admin",
             generate_password_hash("Ajmer@2026")),
        )
    con.commit()
    con.close()

init()

def current():
    uid = session.get("uid")
    if not uid:
        return None
    con = db()
    row = con.execute("SELECT * FROM users WHERE id=? AND active=1", (uid,)).fetchone()
    con.close()
    return row

def login_required(fn):
    @wraps(fn)
    def wrap(*a, **k):
        if not current():
            return jsonify({"error": "लॉगिन जरूरी है"}), 401
        return fn(*a, **k)
    return wrap

def admin_required(fn):
    @wraps(fn)
    def wrap(*a, **k):
        u = current()
        if not u or not can_see_all(u):
            return jsonify({"error": "सिर्फ System Admin"}), 403
        return fn(*a, **k)
    return wrap

def send_sms(mobile, code):
    key = os.environ.get("FAST2SMS_API_KEY", "").strip()
    if not key:
        return False
    phone = "".join(ch for ch in mobile if ch.isdigit())[-10:]
    body = urllib.parse.urlencode({
        "authorization": key,
        "route": "otp",
        "variables_values": code,
        "numbers": phone,
        "flash": "0",
    }).encode()
    req = urllib.request.Request("https://www.fast2sms.com/dev/bulkV2", data=body, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=20) as res:
            return res.status == 200
    except Exception:
        return False

def send_text(mobile, text):
    key = os.environ.get("FAST2SMS_API_KEY", "").strip()
    if not key:
        return False
    phone = "".join(ch for ch in mobile if ch.isdigit())[-10:]
    body = urllib.parse.urlencode({
        "authorization": key,
        "route": "q",
        "message": text,
        "numbers": phone,
        "flash": "0",
    }).encode()
    req = urllib.request.Request("https://www.fast2sms.com/dev/bulkV2", data=body, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=20) as res:
            return res.status == 200
    except Exception:
        return False


def manager_required(fn):
    @wraps(fn)
    def wrap(*a, **k):
        u = current()
        if not u or u["role"] not in ("admin", "district_admin"):
            return jsonify({"error": "System Admin या District Admin"}), 403
        return fn(*a, **k)
    return wrap


def audit(con, user, action, entity, entity_id="", detail=""):
    con.execute("CREATE TABLE IF NOT EXISTS audit_logs (id TEXT PRIMARY KEY, user_name TEXT, role TEXT, action TEXT, entity TEXT, entity_id TEXT, detail TEXT, ip TEXT, at TEXT)")
    con.execute(
        "INSERT INTO audit_logs VALUES (?,?,?,?,?,?,?,?,?)",
        (uuid.uuid4().hex, (user or {}).get("name"), (user or {}).get("role"), action, entity, entity_id, detail, request.remote_addr or "", datetime.now().isoformat(timespec="seconds")),
    )

def add_history(con, rid, status, note, by):
    con.execute(
        "INSERT INTO history VALUES (?,?,?,?,?,?)",
        (uuid.uuid4().hex, rid, status, note, by, datetime.now().isoformat(timespec="seconds")),
    )

@app.get("/")
def home():
    return send_from_directory(app.static_folder, "index.html")

@app.post("/api/login")
def login():
    ip = request.remote_addr or "local"
    if FAILS.get(ip, 0) >= 8:
        return jsonify({"error": "Too many failed logins. Try again later."}), 429
    data = request.json or {}
    con = db()
    user = con.execute("SELECT * FROM users WHERE username=? AND active=1", (data.get("username", "").strip(),)).fetchone()
    con.close()
    if not user or not check_password_hash(user["password_hash"], data.get("password", "")):
        return jsonify({"error": "गलत यूजरनेम या पासवर्ड"}), 401
    FAILS[ip] = 0
    session.permanent = True
    session["uid"] = user["id"]
    con = db()
    audit(con, dict(user), "login", "session", user["id"])
    con.commit(); con.close()
    return jsonify({"id": user["id"], "name": user["name"], "role": user["role"], "cluster": user["cluster"], "username": user["username"], "mobile": user["mobile"]})

@app.post("/api/password")
@login_required
def change_password():
    u = current()
    data = request.json or {}
    pw = data.get("password") or ""
    if len(pw) < 8 or pw.lower()==pw or pw.upper()==pw or not any(ch.isdigit() for ch in pw):
        return jsonify({"error": "Password needs 8 characters, upper, lower, and a number"}), 400
    if not check_password_hash(u["password_hash"], data.get("old") or ""):
        return jsonify({"error": "Current password is wrong"}), 400
    con = db()
    con.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(pw), u["id"]))
    audit(con, u, "password_change", "user", u["id"])
    con.commit(); con.close()
    return jsonify({"ok": True})

@app.get("/api/backup")
@admin_required
def backup():
    from flask import send_file
    return send_file(DB, as_attachment=True, download_name="ajmer-krishi-backup.db")

@app.get("/api/notifications")
@login_required
def notifications():
    u = current()
    con = db()
    rows = [dict(r) for r in con.execute("SELECT status, note, by_name, at FROM history ORDER BY at DESC LIMIT 30").fetchall()]
    con.close()
    return jsonify(rows)

@app.post("/api/logout")
def logout():
    session.clear()
    return jsonify({"ok": True})

@app.get("/api/me")
def me():
    u = current()
    if not u:
        return jsonify(None)
    return jsonify({"id": u["id"], "name": u["name"], "role": u["role"], "cluster": u["cluster"], "username": u["username"], "mobile": u["mobile"]})

TEHSILS = ["अजमेर","पीसांगन","पुष्कर","किशनगढ़","नसीराबाद","श्रीनगर","अराई","केकड़ी","भिनाय","सरवाड़","सावर"]

def lgd_tree():
    path = os.path.join(BASE, "ajmer_lgd.json")
    tree = json.loads(open(path, encoding="utf-8").read()) if os.path.exists(path) else {}
    con = db()
    con.execute("CREATE TABLE IF NOT EXISTS extra_villages (id TEXT PRIMARY KEY, tehsil TEXT, gp TEXT, village TEXT)")
    for row in con.execute("SELECT tehsil, gp, village FROM extra_villages"):
        tree.setdefault(row["tehsil"], {}).setdefault(row["gp"] or "Unassigned", [])
        if row["village"] not in tree[row["tehsil"]][row["gp"] or "Unassigned"]:
            tree[row["tehsil"]][row["gp"] or "Unassigned"].append(row["village"])
    con.close()
    return tree

@app.get("/api/directory")
def directory():
    try:
        return jsonify({"district": "अजमेर", "tree": lgd_tree()})
    except Exception:
        return jsonify({"district": "अजमेर", "tree": {}})

@app.post("/api/directory/village")
@admin_required

def cluster_for(con, scheme, gp):
    con.execute("CREATE TABLE IF NOT EXISTS scheme_clusters (id TEXT PRIMARY KEY, scheme TEXT, cluster_id TEXT, gp TEXT)")
    row = con.execute("SELECT c.name FROM scheme_clusters s JOIN clusters c ON c.id=s.cluster_id WHERE s.scheme=? AND lower(s.gp)=lower(?)", (scheme, gp or "")).fetchone()
    return row["name"] if row else ""

def attach_farmers(con, scheme, gp, name):
    if not gp or not name:
        return
    for r in con.execute("SELECT id, payload FROM records WHERE form_type=?", (scheme,)).fetchall():
        item = json.loads(r["payload"] or "{}")
        if (item.get("GP") or "").strip().lower() != gp.strip().lower():
            continue
        item["Cluster ID"] = name
        con.execute("UPDATE records SET cluster=?, payload=? WHERE id=?", (name, json.dumps(item, ensure_ascii=False), r["id"]))

def map_village():
    data = request.json or {}
    tehsil, gp, village = (data.get("tehsil") or "").strip(), (data.get("gp") or "").strip(), (data.get("village") or "").strip()
    if not tehsil or not village:
        return jsonify({"error": "तहसील और गाँव जरूरी"}), 400
    con = db()
    con.execute("CREATE TABLE IF NOT EXISTS extra_villages (id TEXT PRIMARY KEY, tehsil TEXT, gp TEXT, village TEXT)")
    con.execute("INSERT INTO extra_villages VALUES (?,?,?,?)", (uuid.uuid4().hex, tehsil, gp or "Unassigned", village))
    con.commit()
    con.close()
    return jsonify({"ok": True})

@app.post("/api/map/cluster")
@manager_required
def map_cluster():
    data = request.json or {}
    scheme = data.get("scheme")
    if scheme not in ("pkvy", "natural", "minikit", "demo"):
        return jsonify({"error": "Choose PKVY, Natural Farming, Minikit, or Demonstration"}), 400
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Cluster name is required"}), 400
    con = db()
    row = con.execute("SELECT id FROM clusters WHERE name=?", (name,)).fetchone()
    cid = row["id"] if row else uuid.uuid4().hex
    code = (scheme or "CL")[:3].upper() + "-2026-" + cid[:4].upper()
    if not row:
        con.execute("INSERT INTO clusters VALUES (?,?,?,?)", (cid, "अजमेर", data.get("tehsil") or "", name))
    con.execute("CREATE TABLE IF NOT EXISTS scheme_clusters (id TEXT PRIMARY KEY, scheme TEXT, cluster_id TEXT, gp TEXT)")
    con.execute("INSERT INTO scheme_clusters VALUES (?,?,?,?)", (uuid.uuid4().hex, scheme, cid, data.get("gp") or ""))
    for village in data.get("villages") or []:
        con.execute("INSERT INTO places VALUES (?,?,?,?)", (uuid.uuid4().hex, cid, village, data.get("gp") or village))
    attach_farmers(con, scheme, data.get("gp") or "", name)
    con.commit()
    con.close()
    return jsonify({"ok": True, "id": cid, "cluster_id": code})

@app.post("/api/map/gp")
@manager_required
def map_gp():
    data = request.json or {}
    username = (data.get("username") or "").strip()
    gp = (data.get("gp") or "").strip()
    tehsil = (data.get("tehsil") or "").strip()
    con = db()
    user = con.execute("SELECT * FROM users WHERE username=? AND role='supervisor' AND active=1", (username,)).fetchone()
    if not user:
        con.close()
        return jsonify({"error": "Agriculture Supervisor नहीं मिला"}), 404
    con.execute("CREATE TABLE IF NOT EXISTS supervisor_gp (id TEXT PRIMARY KEY, user_id TEXT, tehsil TEXT, gp TEXT)")
    con.execute("INSERT INTO supervisor_gp VALUES (?,?,?,?)", (uuid.uuid4().hex, user["id"], tehsil, gp))
    cname = gp + " cluster"
    row = con.execute("SELECT id FROM clusters WHERE name=?", (cname,)).fetchone()
    cid = row["id"] if row else uuid.uuid4().hex
    if not row:
        con.execute("INSERT INTO clusters VALUES (?,?,?,?)", (cid, "अजमेर", tehsil, cname))
    try:
        tree = lgd_tree()
        villages = ((tree.get(tehsil) or {}).get(gp) or []) if isinstance(tree.get(tehsil), dict) else []
        for village in villages:
            exists = con.execute("SELECT 1 FROM places WHERE cluster_id=? AND village=?", (cid, village)).fetchone()
            if not exists:
                con.execute("INSERT INTO places VALUES (?,?,?,?)", (uuid.uuid4().hex, cid, village, gp))
        con.execute("UPDATE users SET cluster=? WHERE id=?", (cname, user["id"]))
        con.commit()
    except Exception as exc:
        con.close()
        return jsonify({"error": "Supervisor map failed: " + str(exc)}), 400
    con.close()
    return jsonify({"ok": True, "cluster": cname})

@app.post("/api/map/aao")
@manager_required
def map_aao():
    data = request.json or {}
    con = db()
    user = con.execute("SELECT * FROM users WHERE username=? AND role='aao'", (data.get("username"),)).fetchone()
    if not user:
        con.close()
        return jsonify({"error": "AAO user not found"}), 404
    con.execute("CREATE TABLE IF NOT EXISTS aao_map (id TEXT PRIMARY KEY, user_id TEXT, block TEXT, office_type TEXT, office_name TEXT, supervisors TEXT)")
    con.execute("INSERT INTO aao_map VALUES (?,?,?,?,?,?)", (uuid.uuid4().hex, user["id"], data.get("block") or "", data.get("office_type") or "block", data.get("office_name") or "", ",".join(data.get("supervisors") or [])))
    con.commit(); con.close()
    return jsonify({"ok": True})

@app.get("/api/meta")
def meta():
    con = db()
    out = {}
    con.execute("CREATE TABLE IF NOT EXISTS scheme_clusters (id TEXT PRIMARY KEY, scheme TEXT, cluster_id TEXT, gp TEXT)")
    for c in con.execute("SELECT * FROM clusters ORDER BY name").fetchall():
        places = [dict(p) for p in con.execute("SELECT id, village, gram_panchayat FROM places WHERE cluster_id=?", (c["id"],)).fetchall()]
        schemes = [r["scheme"] for r in con.execute("SELECT scheme FROM scheme_clusters WHERE cluster_id=?", (c["id"],)).fetchall()]
        out[c["name"]] = {"id": c["id"], "district": c["district"], "block": c["block"], "places": places, "schemes": schemes, "gp": places[0]["gram_panchayat"] if places else ""}
    con.close()
    return jsonify(out)

@app.post("/api/clusters")
@manager_required
def add_cluster():
    data = request.json or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "क्लस्टर नाम लिखें"}), 400
    con = db()
    try:
        con.execute("INSERT INTO clusters VALUES (?,?,?,?)", (uuid.uuid4().hex, data.get("district") or "अजमेर", data.get("block") or name, name))
        con.commit()
    except sqlite3.IntegrityError:
        con.close()
        return jsonify({"error": "यह क्लस्टर पहले से है"}), 400
    con.close()
    return jsonify({"ok": True})

@app.post("/api/farmers/by-gp")
@admin_required
def delete_farmers_gp():
    data = request.json or {}
    gp = (data.get("gp") or "").strip().lower()
    scheme = data.get("scheme") or ""
    if not gp:
        return jsonify({"error": "Gram Panchayat is required"}), 400
    con = db()
    removed = 0
    rows = con.execute("SELECT id, payload, form_type FROM records").fetchall()
    for r in rows:
        item = json.loads(r["payload"] or "{}")
        if scheme and (item.get("scheme") or r["form_type"]) != scheme:
            continue
        if (item.get("GP") or "").strip().lower() == gp:
            con.execute("DELETE FROM records WHERE id=?", (r["id"],))
            removed += 1
    con.commit(); con.close()
    return jsonify({"ok": True, "removed": removed})

@app.delete("/api/clusters/<cid>")
@admin_required
def del_cluster(cid):
    con = db()
    con.execute("DELETE FROM places WHERE cluster_id=?", (cid,))
    con.execute("DELETE FROM clusters WHERE id=?", (cid,))
    con.commit()
    con.close()
    return jsonify({"ok": True})

@app.post("/api/places")
@manager_required
def add_place():
    data = request.json or {}
    if not data.get("cluster_id") or not data.get("village"):
        return jsonify({"error": "क्लस्टर और गाँव जरूरी"}), 400
    con = db()
    con.execute("INSERT INTO places VALUES (?,?,?,?)", (uuid.uuid4().hex, data["cluster_id"], data["village"].strip(), (data.get("gram_panchayat") or data["village"]).strip()))
    con.commit()
    con.close()
    return jsonify({"ok": True})

@app.delete("/api/places/<pid>")
@admin_required
def del_place(pid):
    con = db()
    con.execute("DELETE FROM places WHERE id=?", (pid,))
    con.commit()
    con.close()
    return jsonify({"ok": True})

@app.post("/api/otp/send")
def otp_send():
    data = request.json or {}
    username = (data.get("username") or "").strip()
    mobile = (data.get("mobile") or "").strip()
    con = db()
    user = con.execute("SELECT * FROM users WHERE username=? AND active=1", (username,)).fetchone()
    if not user or (user["mobile"] or "")[-10:] != mobile[-10:]:
        con.close()
        return jsonify({"error": "यूजरनेम और मोबाइल मेल नहीं खाते"}), 400
    code = f"{random.randint(100000, 999999)}"
    exp = (datetime.now() + timedelta(minutes=10)).isoformat(timespec="seconds")
    con.execute("INSERT INTO otps VALUES (?,?,?,?,?,0)", (uuid.uuid4().hex, username, mobile, code, exp))
    con.commit()
    con.close()
    sent = send_sms(mobile, code)
    if sent:
        return jsonify({"ok": True, "message": "OTP मोबाइल पर भेज दिया। 10 मिनट तक मान्य है।"})
    return jsonify({"ok": True, "message": "OTP बन गया, SMS कुंजी अभी सेट नहीं है। एडमिन Users में OTP देख सकता है।"})

@app.post("/api/otp/reset")
def otp_reset():
    data = request.json or {}
    if len(data.get("password") or "") < 8:
        return jsonify({"error": "नया Password must be at least 8 characters"}), 400
    con = db()
    row = con.execute("SELECT * FROM otps WHERE username=? AND code=? AND used=0 ORDER BY expires DESC", (data.get("username"), data.get("otp"))).fetchone()
    if not row or row["expires"] < datetime.now().isoformat(timespec="seconds"):
        con.close()
        return jsonify({"error": "OTP गलत या खत्म हो गया"}), 400
    con.execute("UPDATE users SET password_hash=? WHERE username=?", (generate_password_hash(data["password"]), data.get("username")))
    con.execute("UPDATE otps SET used=1 WHERE id=?", (row["id"],))
    con.commit()
    con.close()
    return jsonify({"ok": True})

@app.get("/api/otps")
@admin_required
def list_otps():
    con = db()
    rows = [dict(r) for r in con.execute("SELECT username, mobile, code, expires, used FROM otps ORDER BY expires DESC LIMIT 20").fetchall()]
    con.close()
    return jsonify(rows)

@app.post("/api/reset-request")
def reset_request():
    username = (request.json or {}).get("username", "").strip()
    if not username:
        return jsonify({"error": "यूजरनेम लिखें"}), 400
    con = db()
    con.execute("INSERT INTO resets VALUES (?,?,?,?)", (uuid.uuid4().hex, username, "pending", datetime.now().isoformat(timespec="seconds")))
    con.commit()
    con.close()
    return jsonify({"ok": True, "message": "अनुरोध एडमिन के पास चला गया"})

@app.get("/api/resets")
@admin_required
def resets():
    con = db()
    rows = [dict(r) for r in con.execute("SELECT * FROM resets ORDER BY created_at DESC").fetchall()]
    con.close()
    return jsonify(rows)

@app.post("/api/admin/reset")
@manager_required
def admin_reset():
    data = request.json or {}
    password = data.get("password", "")
    if len(password) < 8:
        return jsonify({"error": "Password must be at least 8 characters"}), 400
    con = db()
    con.execute("UPDATE users SET password_hash=? WHERE username=?", (generate_password_hash(password), data.get("username")))
    con.execute("UPDATE resets SET status='done' WHERE username=?", (data.get("username"),))
    con.commit()
    con.close()
    return jsonify({"ok": True})

@app.post("/api/users")
@manager_required
def add_user():
    data = request.json or {}
    me = current()
    allowed = ("aao", "supervisor", "lrp", "krishi_sakhi", "crp") if me["role"] == "district_admin" else ("district_admin", "aao", "supervisor", "lrp", "krishi_sakhi", "crp")
    if data.get("role") not in allowed:
        return jsonify({"error": "District Admin can only create cluster users"}), 403
    pw = data.get("password", "")
    if len(pw) < 8 or pw.lower()==pw or pw.upper()==pw or not any(ch.isdigit() for ch in pw):
        return jsonify({"error": "Password must be at least 8 characters"}), 400
    mobile = (data.get("mobile") or "").strip()
    if len(mobile) < 10:
        return jsonify({"error": "Mobile number is required"}), 400
    con = db()
    try:
        con.execute(
            "INSERT INTO users VALUES (?,?,?,?,?,?,?,1)",
            (uuid.uuid4().hex, data["role"], data.get("name"), data.get("cluster") or "All",
             mobile, data["username"].strip(), generate_password_hash(data["password"])),
        )
        con.commit()
    except sqlite3.IntegrityError:
        con.close()
        return jsonify({"error": "Username already exists"}), 400
    con.close()
    link = os.environ.get("SITE_URL", "https://ajmer-krishi.onrender.com")
    text = f"Ajmer Agriculture Information\nLink: {link}\nUsername: {data['username'].strip()}\nPassword: {data['password']}\nCluster: {data.get('cluster') or 'All'}"
    sent = send_text(mobile, text)
    return jsonify({"ok": True, "sms": sent})

@app.get("/api/users")
@manager_required
def users():
    con = db()
    rows = [dict(r) for r in con.execute("SELECT id, role, name, cluster, mobile, username, active FROM users").fetchall()]
    con.close()
    return jsonify(rows)

@app.delete("/api/users/<uid>")
@admin_required
def delete_user(uid):
    me = current()
    if uid == me["id"]:
        return jsonify({"error": "अपना लॉगिन नहीं मिटा सकते"}), 400
    con = db()
    con.execute("UPDATE users SET active=0 WHERE id=?", (uid,))
    con.commit()
    con.close()
    return jsonify({"ok": True})

SCHEME_BY_ROLE = {
    "lrp": {"pkvy"},
    "krishi_sakhi": {"natural"},
    "crp": {"natural"},
    "supervisor": {"pkvy", "natural", "minikit", "demo"},
}

def can_see_all(u):
    return u["role"] in ("admin", "district_admin", "supervisor", "aao")

def visible_sql(u):
    if u["role"] in ("admin", "district_admin"):
        return "", []
    return " WHERE cluster=?", [u["cluster"]]

@app.get("/api/records")
@login_required
def records():
    u = current()
    where, args = visible_sql(u)
    status = request.args.get("status")
    if status:
        where = (where + " AND " if where else " WHERE ") + "status=?"
        args.append(status)
    con = db()
    rows = [dict(r) for r in con.execute(f"SELECT * FROM records{where} ORDER BY created_at DESC", args).fetchall()]
    for r in rows:
        r["payload"] = json.loads(r["payload"] or "{}")
        r["photos"] = [dict(p) for p in con.execute("SELECT id, filename FROM photos WHERE record_id=?", (r["id"],)).fetchall()]
        r["history"] = [dict(h) for h in con.execute("SELECT status, note, by_name, at FROM history WHERE record_id=? ORDER BY at", (r["id"],)).fetchall()]
        con.execute("CREATE TABLE IF NOT EXISTS signatures (id TEXT PRIMARY KEY, record_id TEXT, role TEXT, by_name TEXT, image TEXT, at TEXT)")
        r["signatures"] = [dict(x) for x in con.execute("SELECT role, by_name, image, at FROM signatures WHERE record_id=?", (r["id"],)).fetchall()]
    allowed = SCHEME_BY_ROLE.get(u["role"])
    if allowed:
        rows = [r for r in rows if (r["payload"].get("scheme") in allowed) or r["form_type"] in allowed or r["form_type"] == "farmer_master"]
    con.close()
    return jsonify(rows)

@app.post("/api/records")
@login_required
def create_record():
    u = current()
    data = request.json or {}
    cluster = u["cluster"] if u["role"] not in ("admin", "district_admin") else (data.get("cluster") or u["cluster"])
    if u["role"] not in ("admin", "district_admin") and data.get("cluster") and data.get("cluster") != u["cluster"]:
        return jsonify({"error": "आप केवल अपने क्लस्टर का डेटा भर सकते हैं"}), 403
    scheme = (data.get("form_type") or "")
    allowed = SCHEME_BY_ROLE.get(u["role"])
    if allowed and scheme not in allowed and scheme != "farmer_master":
        return jsonify({"error": "This scheme is not allowed for your role"}), 403
    rid = uuid.uuid4().hex
    now = datetime.now().isoformat(timespec="seconds")
    con = db()
    con.execute(
        "INSERT INTO records VALUES (?,?,?,?,?,?,?,?,?,?)",
        (rid, u["id"], u["name"], u["role"], cluster, data.get("village", ""), data.get("form_type", "परिशिष्ट 8"),
         json.dumps(data.get("payload") or {}, ensure_ascii=False), "submitted", now),
    )
    add_history(con, rid, "submitted", "जमा किया", u["name"])
    con.commit()
    con.close()
    return jsonify({"id": rid})

@app.post("/api/records/<rid>/photo")
@login_required
def upload_photo(rid):
    u = current()
    con = db()
    rec = con.execute("SELECT * FROM records WHERE id=?", (rid,)).fetchone()
    if not rec or (not can_see_all(u) and rec["cluster"] != u["cluster"]):
        con.close()
        return jsonify({"error": "रिपोर्ट नहीं मिली"}), 404
    f = request.files.get("photo")
    if not f:
        con.close()
        return jsonify({"error": "फोटो नहीं"}), 400
    ext = os.path.splitext(secure_filename(f.filename or "photo.jpg"))[1] or ".jpg"
    name = rid + "_" + uuid.uuid4().hex[:8] + ext
    f.save(os.path.join(UPLOAD, name))
    pid = uuid.uuid4().hex
    con.execute("INSERT INTO photos VALUES (?,?,?,?)", (pid, rid, name, datetime.now().isoformat(timespec="seconds")))
    con.commit()
    con.close()
    return jsonify({"id": pid})

@app.get("/api/photos/<pid>")
@login_required
def photo(pid):
    u = current()
    con = db()
    row = con.execute(
        "SELECT photos.filename, records.cluster FROM photos JOIN records ON records.id=photos.record_id WHERE photos.id=?",
        (pid,),
    ).fetchone()
    con.close()
    if not row or (not can_see_all(u) and row["cluster"] != u["cluster"]):
        return jsonify({"error": "नहीं"}), 404
    return send_from_directory(UPLOAD, row["filename"])

@app.post("/api/records/<rid>/edit")
@login_required
def edit_record(rid):
    u = current()
    data = request.json or {}
    con = db()
    rec = con.execute("SELECT * FROM records WHERE id=?", (rid,)).fetchone()
    if not rec:
        con.close()
        return jsonify({"error": "Record not found"}), 404
    if rec["status"] not in ("submitted", "draft", "returned", "correction_required"):
        con.close()
        return jsonify({"error": "Signed record cannot be edited. Ask District Admin to revert it."}), 400
    payload = json.loads(rec["payload"] or "{}")
    payload.update(data.get("payload") or {})
    new_status = "submitted" if rec["status"] in ("returned", "correction_required", "corrected") else rec["status"]
    con.execute("UPDATE records SET payload=?, village=?, status=? WHERE id=?", (json.dumps(payload, ensure_ascii=False), payload.get("Village") or rec["village"], new_status, rid))
    add_history(con, rid, "corrected", "Edited before sign", u["name"])
    con.commit(); con.close()
    return jsonify({"ok": True})

@app.post("/api/records/<rid>/status")
@admin_required
def set_status(rid):
    data = request.json or {}
    status = data.get("status")
    if status not in ("submitted", "correction_required", "resubmitted", "approved"):
        return jsonify({"error": "स्टेटस गलत"}), 400
    u = current()
    con = db()
    con.execute("UPDATE records SET status=? WHERE id=?", (status, rid))
    add_history(con, rid, status, data.get("note", ""), u["name"])
    con.commit()
    con.close()
    return jsonify({"ok": True})

@app.post("/api/records/<rid>/resubmit")
@login_required
def resubmit(rid):
    u = current()
    con = db()
    rec = con.execute("SELECT * FROM records WHERE id=?", (rid,)).fetchone()
    if not rec or rec["user_id"] != u["id"]:
        con.close()
        return jsonify({"error": "नहीं मिली"}), 404
    payload = request.json or {}
    con.execute("UPDATE records SET payload=?, status='resubmitted' WHERE id=?", (json.dumps(payload.get("payload") or {}, ensure_ascii=False), rid))
    add_history(con, rid, "resubmitted", payload.get("note", "दोबारा जमा"), u["name"])
    con.commit()
    con.close()
    return jsonify({"ok": True})

REQUIRED = ["farmer", "crop", "kind", "area", "yield"]

DEV = {"अ":"a","आ":"a","इ":"i","ई":"i","उ":"u","ऊ":"u","ए":"e","ऐ":"ai","ओ":"o","औ":"au","क":"k","ख":"kh","ग":"g","घ":"gh","च":"ch","छ":"chh","ज":"j","झ":"jh","ट":"t","ठ":"th","ड":"d","ढ":"dh","त":"t","थ":"th","द":"d","ध":"dh","न":"n","प":"p","फ":"ph","ब":"b","भ":"bh","म":"m","य":"y","र":"r","ल":"l","व":"v","श":"sh","ष":"sh","स":"s","ह":"h","ा":"a","ि":"i","ी":"i","ु":"u","ू":"u","े":"e","ै":"ai","ो":"o","ौ":"au","ं":"n","ँ":"n","्":""}
def to_english(value):
    text = str(value or "")
    if not any("\u0900" <= ch <= "\u097F" for ch in text):
        return text
    return "".join(DEV.get(ch, ch) for ch in text)

def norm(v):
    return str(v or "").strip().lower()


TEMPLATES = {
  "pkvy": ["Farmer Name","Father/Husband Name","Jan Aadhaar","Mobile","Aadhaar No","District","Tehsil","Block","GP","Village","Financial Year","Khata No/Plot No","Khasra No","Total Area","Offered Area","Crop","Bank Account No","IFSC Code","Branch Address","Cow","Buffalo","Goat","Irrigation Source","Land Type","Last Date of Prohibited Input"],
  "natural": ["Farmer Name","Father/Husband Name","Jan Aadhaar","Mobile","Aadhaar No","District","Tehsil","Block","GP","Village","Financial Year","Khata No/Plot No","Khasra No","Total Area","Offered Area","Crop","Bank Account No","IFSC Code","Branch Address","Cow","Buffalo","Goat","Irrigation Source","Land Type","Last Date of Prohibited Input"],
  "minikit": ["Farmer ID","Farmer Name","Father/Husband Name","Jan Aadhaar","Mobile","District","Tehsil","Block","GP","Village","Cluster ID","Financial Year","Demonstration Type","Demonstration Category","Season","Demonstration Date","Khasra No.","Area","Crop","Variety"],
  "demo": ["Farmer ID","Farmer Name","Father/Husband Name","Jan Aadhaar","Mobile","District","Tehsil","Block","GP","Village","Cluster ID","Financial Year","Demonstration Type","Demonstration Category","Season","Demonstration Date","Khasra No.","Area","Crop","Variety"],
}
REQUIRED_BY = {k: ["Farmer Name","Village","Mobile"] for k in TEMPLATES}

@app.get("/api/template")
@login_required
def template():
    scheme = request.args.get("scheme") or "pkvy"
    cols = TEMPLATES.get(scheme)
    if not cols:
        return jsonify({"error": "योजना नहीं मिली"}), 400
    wb = Workbook()
    ws = wb.active
    ws.title = scheme
    ws.append(cols)
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    from flask import send_file
    return send_file(bio, as_attachment=True, download_name=f"{scheme}-template.xlsx")

@app.post("/api/excel")
@login_required
def excel_upload():
    u = current()
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "File is missing"}), 400
    try:
        wb = load_workbook(f, data_only=True)
    except Exception:
        return jsonify({"error": "This is not a valid Excel file. Download the template and upload .xlsx"}), 400
    ws = wb.active
    scheme = request.form.get("scheme") or "pkvy"
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(max_row=1))]
    missing_cols = [c for c in TEMPLATES.get(scheme, []) if c not in headers]
    if missing_cols:
        err_id = uuid.uuid4().hex
        ew = Workbook(); ews = ew.active
        ews.append(["Row", "Field", "Entered", "Error", "Correction"])
        ews.append([1, "Header", ", ".join(headers), "कॉलम नहीं मिले", "टेम्पलेट वाले कॉलम रखें: " + ", ".join(missing_cols)])
        os.makedirs(os.path.join(BASE, "data"), exist_ok=True)
        ew.save(os.path.join(BASE, "data", err_id + ".xlsx"))
        return jsonify({"saved": 0, "errors": [{"row": 1, "type": "error", "detail": "Missing columns: " + ", ".join(missing_cols)}], "error_file": err_id, "message": "Excel header does not match the template"})
    def col(*parts):
        for i, h in enumerate(headers):
            if any(p in h.lower() for p in parts):
                return i
        return None
    idx = {
        "farmer": col("किसान", "farmer"),
        "crop": col("फसल", "crop"),
        "kind": col("किस्म", "kind"),
        "area": col("क्षेत्र", "area"),
        "yield": col("उपज", "yield"),
        "ytype": col("अनुमानित", "वास्तविक", "type"),
    }
    con = db()
    seen = set()
    for r in con.execute("SELECT payload FROM records").fetchall():
        old = json.loads(r["payload"])
        seen.add(norm(old.get("Jan Aadhaar")) + "|" + norm(old.get("Aadhaar No")))
    clean, errors = [], []
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        item = {}
        for i, h in enumerate(headers):
            val = row[i] if i < len(row) else ""
            item[h] = "" if val is None else to_english(val)
        item["scheme"] = scheme
        fid = str(item.get("Farmer ID") or "")
        prefix = {"pkvy":"PKVY-F","natural":"NF-F","minikit":"MK-F","demo":"DEM-F"}.get(scheme)
        if fid and prefix and not fid.startswith(prefix):
            errors.append({"row": n, "field": "Farmer ID", "entered": fid, "type": "error", "detail": "This farmer does not belong to this scheme"})
            continue
        item["farmer"] = item.get("Farmer Name") or item.get("farmer")
        item["crop"] = item.get("Crop") or item.get("crop")
        if scheme in ("pkvy", "natural"):
            prefix = {"pkvy":"PKVY-F","natural":"NF-F"}[scheme]
            item["Farmer ID"] = item["farmer_id"] = prefix + "-" + uuid.uuid4().hex[:6].upper()
            gp_name=(item.get("GP") or "").strip(); village=(item.get("Village") or "").strip(); item["Cluster ID"] = cluster_for(con, scheme, gp_name) or (gp_name if gp_name.lower()==village.lower() else (gp_name+"-"+village).strip("-"))
            source = (item.get("Irrigation Source") or "").strip().lower()
            item["IRRIGATED OR NON IRRIGATED"] = "NON IRRIGATED" if source in ("", "no", "none") else "IRRIGATED"
            item["Branch Address"] = item.get("GP") or item.get("Block") or ""
            offered = float(item.get("Offered Area") or 0)
            total = float(item.get("Total Area") or 0)
            if offered and total and total < offered:
                errors.append({"row": n, "field": "Total Area", "entered": item.get("Total Area"), "type": "error", "detail": "Total area cannot be less than offered area"})
                continue
            ifsc = (item.get("IFSC Code") or "").strip().upper()
            if ifsc and not (len(ifsc)==11 and ifsc[:4].isalpha() and ifsc[4]=="0"):
                errors.append({"row": n, "field": "IFSC Code", "entered": ifsc, "type": "error", "detail": "Invalid IFSC code"})
                continue
            year = (item.get("Financial Year") or "").strip()
            if u["role"] not in ("admin", "district_admin") and year and year != "2026-27":
                errors.append({"row": n, "field": "Financial Year", "entered": year, "type": "error", "detail": "Only current year is allowed. Old year can be added by System Admin or District Admin"})
                continue
            key = norm(item.get("Farmer Name")) + "|" + norm(item.get("Jan Aadhaar"))
            acc = norm(item.get("Bank Account No"))
            if acc and acc in seen:
                errors.append({"row": n, "field": "Bank Account No", "entered": item.get("Bank Account No"), "type": "duplicate", "detail": "Duplicate bank account number"})
                continue
            seen.add(key); seen.add(acc)
        missing = [k for k in REQUIRED_BY.get(scheme, []) if str(item.get(k) or "").strip() == ""]
        mobile = "".join(ch for ch in str(item.get("Mobile") or "") if ch.isdigit())
        jan = "".join(ch for ch in str(item.get("Jan Aadhaar") or "") if ch.isdigit())
        aad = "".join(ch for ch in str(item.get("Aadhaar No") or "") if ch.isdigit())
        if mobile and len(mobile) != 10:
            errors.append({"row": n, "field": "Mobile", "entered": item.get("Mobile"), "type": "error", "detail": "Mobile must be 10 digits"})
            continue
        if jan and len(jan) != 10:
            errors.append({"row": n, "field": "Jan Aadhaar", "entered": item.get("Jan Aadhaar"), "type": "error", "detail": "Jan Aadhaar must be 10 digits"})
            continue
        if aad and len(aad) != 12:
            errors.append({"row": n, "field": "Aadhaar No", "entered": item.get("Aadhaar No"), "type": "error", "detail": "Aadhaar must be 12 digits"})
            continue
        idkey = norm(item.get("Jan Aadhaar")) + "|" + norm(item.get("Aadhaar No"))
        if missing:
            errors.append({"row": n, "field": ", ".join(missing), "entered": item.get("Farmer Name") or "", "type": "error", "detail": ", ".join(missing) + " is empty or invalid"})
            continue
        if idkey.strip("|") and idkey in seen:
            errors.append({"row": n, "field": "Aadhaar No", "entered": item.get("Aadhaar No") or "", "type": "duplicate", "detail": "Duplicate only when Jan Aadhaar and Aadhaar are both same"})
            continue
        seen.add(idkey)
        clean.append(item)
    if request.form.get("preview") == "1":
        con.close()
        return jsonify({"saved": 0, "preview": clean[:20], "valid": len(clean), "errors": errors, "message": f"Preview: {len(clean)} valid, {len(errors)} invalid"})
    saved = []
    now = datetime.now().isoformat(timespec="seconds")
    for item in clean:
        rid = uuid.uuid4().hex
        con.execute(
            "INSERT INTO records VALUES (?,?,?,?,?,?,?,?,?,?)",
            (rid, u["id"], u["name"], u["role"], item.get("Cluster ID") or u["cluster"], str(item.get("Village") or ""), scheme,
             json.dumps(item, ensure_ascii=False), "submitted", now),
        )
        add_history(con, rid, "submitted", "Excel से जमा", u["name"])
        saved.append(rid)
    con.commit()
    con.close()
    err_id = None
    if errors:
        ew = Workbook()
        ews = ew.active
        ews.append(["Row", "Field", "Entered", "Error", "Correction"])
        for e in errors:
            ews.append([e.get("row"), e.get("field") or e.get("detail"), e.get("entered"), e.get("type"), e.get("detail")])
        err_id = uuid.uuid4().hex
        os.makedirs(os.path.join(BASE, "data"), exist_ok=True)
        ew.save(os.path.join(BASE, "data", err_id + ".xlsx"))
    return jsonify({"saved": len(saved), "errors": errors, "error_file": err_id, "message": f"{len(saved)} saved, {len(errors)} error"})

@app.post("/api/photo-excel")
@login_required
def photo_excel():
    f = request.files.get("file")
    scheme = request.form.get("scheme") or "pkvy"
    if not f:
        return jsonify({"ok": False, "reasons": ["Choose a photo"], "rows": []})
    raw = f.read()
    import io, shutil, subprocess, tempfile, re
    from PIL import Image
    path = tempfile.mktemp(suffix=".png")
    try:
        img = Image.open(io.BytesIO(raw)).convert("L")
        img.thumbnail((1600, 1600))
        img.save(path, "PNG")
    except Exception:
        return jsonify({"ok": False, "reasons": ["This file is not a photo"], "rows": []})
    if not shutil.which("tesseract"):
        return jsonify({"ok": False, "reasons": ["Photo reader is not installed on the server"], "rows": []})
    try:
        text = subprocess.check_output(["tesseract", path, "stdout", "-l", "eng"], stderr=subprocess.DEVNULL, text=True, timeout=20)
    except Exception:
        return jsonify({"ok": False, "reasons": ["Photo could not be read. Use a closer table photo."], "rows": []})
    rows = []
    for line in text.splitlines():
        parts = line.split()
        if len(parts) < 6:
            continue
        mobile = next((x for x in parts if re.fullmatch(r"\d{10}", x)), "")
        if not mobile:
            continue
        name = parts[0]
        village = ""
        if "KHERA" in parts:
            village = "NAI KHERA" if "NAI" in parts else "KHERA"
        elif len(parts) >= 2:
            village = parts[-2]
        rows.append({"Farmer Name": name, "Mobile": mobile, "Village": village, "scheme": scheme})
    reasons = []
    if not rows:
        reasons.append("Farmer Name, Mobile and Village were not found in the required template columns")
    return jsonify({"ok": bool(rows), "rows": rows[:50], "text": text[:800], "reasons": reasons, "message": f"{len(rows)} farmers read" if rows else "Data does not match the template"})

@app.post("/api/photo-save")
@login_required
def photo_save():
    u = current()
    data = request.json or {}
    scheme = data.get("scheme") or "pkvy"
    rows = data.get("rows") or []
    if not rows:
        return jsonify({"error": "No rows to submit"}), 400
    con = db()
    now = datetime.now().isoformat(timespec="seconds")
    prefix = {"pkvy": "PKVY-F", "natural": "NF-F", "minikit": "MK-F", "demo": "DEM-F"}.get(scheme, "F")
    for row in rows:
        item = dict(row)
        item["farmer_id"] = prefix + "-" + uuid.uuid4().hex[:6].upper()
        item["farmer"] = row.get("Farmer Name")
        con.execute("INSERT INTO records VALUES (?,?,?,?,?,?,?,?,?,?)", (uuid.uuid4().hex, u["id"], u["name"], u["role"], u["cluster"], row.get("Village") or "", scheme, json.dumps(item, ensure_ascii=False), "submitted", now))
    con.commit(); con.close()
    return jsonify({"ok": True, "saved": len(rows)})

@app.get("/api/errors/<eid>")
@login_required
def error_xlsx(eid):
    path = os.path.join(BASE, "data", eid + ".xlsx")
    if not os.path.exists(path):
        return jsonify({"error": "एरर File is missing"}), 404
    from flask import send_file
    return send_file(path, as_attachment=True, download_name="error-rows.xlsx")

WORKFLOW = ["submitted", "supervisor_signed", "aao_approved"]
NEXT_ROLE = {"submitted": "supervisor", "supervisor_signed": "aao", "signed": "aao"}

@app.post("/api/workflow/<rid>")
@login_required
def workflow(rid):
    u = current()
    data = request.json or {}
    action = data.get("action")
    note = (data.get("note") or "").strip()
    con = db()
    rec = con.execute("SELECT * FROM records WHERE id=?", (rid,)).fetchone()
    if not rec:
        con.close()
        return jsonify({"error": "Record not found"}), 404
    status = rec["status"]
    if action == "return":
        if not note:
            con.close()
            return jsonify({"error": "A return reason is required"}), 400
        new = "returned"
    elif action == "sign":
        actor = data.get("actor") if u["role"] == "admin" and data.get("actor") in ("lrp", "crp", "krishi_sakhi", "supervisor", "aao") else u["role"]
        stage = "submitted" if status in ("returned", "correction_required", "corrected", "draft") else status
        scheme = rec["form_type"]
        chain = ["lrp_signed", "crp_signed", "sakhi_signed", "supervisor_signed", "aao_approved"] if scheme in ("pkvy", "natural") else ["supervisor_signed", "aao_approved"]
        role_status = {"lrp": "lrp_signed", "crp": "crp_signed", "krishi_sakhi": "sakhi_signed", "crp": "sakhi_signed", "supervisor": "supervisor_signed", "aao": "aao_approved"}
        order = ["lrp", "supervisor", "aao"] if scheme == "pkvy" else ["krishi_sakhi", "supervisor", "aao"] if scheme == "natural" else ["supervisor", "aao"]
        step = {"submitted": -1, "lrp_signed": 0, "crp_signed": 1, "sakhi_signed": 2, "supervisor_signed": 3, "signed": 3}.get(stage, -1)
        if scheme not in ("pkvy", "natural") and stage in ("submitted", "draft"):
            step = -1
        expect = order[step + 1] if step + 1 < len(order) else None
        if scheme == "natural" and actor in ("crp", "krishi_sakhi") and stage in ("submitted", "returned", "draft"):
            expect = actor
        if actor != expect and u["role"] != "admin":
            con.close()
            return jsonify({"error": "You cannot sign at this stage"}), 403
        new = role_status.get(actor if u["role"] == "admin" else expect, "signed")
    else:
        con.close()
        return jsonify({"error": "Invalid action"}), 400
    con.execute("UPDATE records SET status=? WHERE id=?", (new, rid))
    add_history(con, rid, new, note or action, u["name"])
    sig = data.get("signature") or ""
    if sig.startswith("data:image"):
        con.execute("CREATE TABLE IF NOT EXISTS signatures (id TEXT PRIMARY KEY, record_id TEXT, role TEXT, by_name TEXT, image TEXT, at TEXT)")
        role = data.get("actor") or u["role"]
        con.execute("DELETE FROM signatures WHERE record_id=? AND role=?", (rid, role))
        con.execute("INSERT INTO signatures VALUES (?,?,?,?,?,?)", (uuid.uuid4().hex, rid, role, u["name"], sig, datetime.now().isoformat(timespec="seconds")))
    try:
        audit(con, u, action, "record", rid, new)
    except Exception:
        pass
    con.commit()
    con.close()
    return jsonify({"ok": True, "status": new})

@app.post("/api/revert")
@login_required
def revert():
    u = current()
    if u["role"] not in ("admin", "district_admin"):
        return jsonify({"error": "Only District Admin can do this"}), 403
    data = request.json or {}
    cid = (data.get("cluster") or "").strip()
    note = (data.get("reason") or "").strip()
    if not cid or not note:
        return jsonify({"error": "Cluster ID and reason are required"}), 400
    con = db()
    rows = con.execute("SELECT id, payload FROM records").fetchall()
    n = 0
    for r in rows:
        pld = json.loads(r["payload"] or "{}")
        if cid in (pld.get("cluster"), pld.get("PKVYClusterID"), pld.get("Cluster ID"), r["id"]):
            con.execute("UPDATE records SET status='returned' WHERE id=?", (r["id"],))
            add_history(con, r["id"], "returned", note, u["name"])
            n += 1
    con.commit()
    con.close()
    return jsonify({"ok": True, "count": n})

@app.get("/api/export")
@login_required
def export():
    u = current()
    where, args = visible_sql(u)
    cluster = request.args.get("cluster")
    if cluster and u["role"] == "admin":
        where, args = " WHERE cluster=?", [cluster]
    con = db()
    rows = con.execute(f"SELECT * FROM records{where} ORDER BY created_at", args).fetchall()
    con.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidated"
    ws.append(["Date", "Cluster", "Village", "Who", "Role", "Status", "Form", "Farmer", "Crop", "Kind", "Area", "Yield"])
    for r in rows:
        p = json.loads(r["payload"] or "{}")
        ws.append([r["created_at"], r["cluster"], r["village"], r["by_name"], r["role"], r["status"], r["form_type"],
                   p.get("farmer") or p.get("farmerName"), p.get("crop"), p.get("kind"), p.get("area") or p.get("areaHa"), p.get("yield") or p.get("yieldQ")])
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    from flask import send_file
    return send_file(bio, as_attachment=True, download_name="ajmer-krishi-consolidated.xlsx")

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 8080)))
